#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script para manejar la persistencia del archivo Excel y asegurar que sea visible.
Este script resuelve el problema de que el archivo Excel se guarde en una ubicación
temporal cuando se ejecuta como .exe compilado.
"""

import os
import sys
import shutil
import pandas as pd
from datetime import datetime
from pathlib import Path

def get_executable_directory():
    """Obtiene el directorio donde se ejecuta el .exe o el script."""
    if getattr(sys, 'frozen', False):
        # Si estamos en un ejecutable PyInstaller
        return os.path.dirname(sys.executable)
    else:
        # Si estamos en desarrollo
        return os.path.dirname(os.path.abspath(__file__))

def ensure_excel_visibility():
    """Asegura que el archivo Excel sea visible y persistente."""
    try:
        # Obtener el directorio del ejecutable
        exe_dir = get_executable_directory()
        excel_file = os.path.join(exe_dir, "unidades_registro.xlsx")
        
        print(f"📁 Directorio del ejecutable: {exe_dir}")
        print(f"📊 Archivo Excel: {excel_file}")
        
        # Verificar si el archivo existe
        if os.path.exists(excel_file):
            print(f"✅ Archivo Excel encontrado: {excel_file}")
            
            # Verificar permisos de escritura
            if os.access(excel_file, os.W_OK):
                print("✅ Archivo Excel tiene permisos de escritura")
            else:
                print("⚠️ Archivo Excel no tiene permisos de escritura")
                
            # Mostrar información del archivo
            file_size = os.path.getsize(excel_file)
            file_time = datetime.fromtimestamp(os.path.getmtime(excel_file))
            print(f"📏 Tamaño: {file_size} bytes")
            print(f"🕰️ Última modificación: {file_time}")
            
            return True
        else:
            print(f"❌ Archivo Excel no encontrado: {excel_file}")
            
            # Intentar crear el archivo si no existe
            print("🔄 Creando archivo Excel...")
            return create_excel_file(exe_dir)
            
    except Exception as e:
        print(f"❌ Error verificando archivo Excel: {e}")
        return False

def create_excel_file(directory):
    """Crea el archivo Excel con la estructura correcta."""
    try:
        excel_file = os.path.join(directory, "unidades_registro.xlsx")
        
        # Crear DataFrame vacío y luego agregar columnas
        df = pd.DataFrame()
        
        # Agregar columnas una por una
        df["🚓 UNIDADES"] = []
        df["🚔 KM"] = []
        df["📌 A.P"] = []
        df["📒 P.O"] = []
        df["🕵️‍♂️ TURNO"] = []
        df["🌙 OBS-TURNO"] = []
        df["🕰️ HORA"] = []
        df["📅 FECHA"] = []
        
        # Guardar con formato
        with pd.ExcelWriter(excel_file, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name='Registro de Unidades')
            
            # Obtener el workbook y worksheet
            workbook = writer.book
            worksheet = writer.sheets['Registro de Unidades']
            
            # Aplicar formato a encabezados
            from openpyxl.styles import Font, Alignment, PatternFill
            
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            
            for col in range(1, 9):  # 8 columnas
                cell = worksheet.cell(row=1, column=col)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center")
        
        print(f"✅ Archivo Excel creado exitosamente: {excel_file}")
        return True
        
    except Exception as e:
        print(f"❌ Error creando archivo Excel: {e}")
        return False

def copy_excel_from_temp():
    """Copia el archivo Excel desde la ubicación temporal si existe."""
    try:
        exe_dir = get_executable_directory()
        target_excel = os.path.join(exe_dir, "unidades_registro.xlsx")
        
        # Buscar el archivo en ubicaciones temporales comunes
        temp_locations = [
            os.path.join(os.environ.get('TEMP', ''), "unidades_registro.xlsx"),
            os.path.join(os.environ.get('TMP', ''), "unidades_registro.xlsx"),
            os.path.join(os.getcwd(), "unidades_registro.xlsx")
        ]
        
        for temp_location in temp_locations:
            if os.path.exists(temp_location):
                print(f"📋 Encontrado archivo en ubicación temporal: {temp_location}")
                
                # Copiar al directorio del ejecutable
                shutil.copy2(temp_location, target_excel)
                print(f"✅ Archivo copiado a: {target_excel}")
                return True
        
        print("ℹ️ No se encontró archivo Excel en ubicaciones temporales")
        return False
        
    except Exception as e:
        print(f"❌ Error copiando archivo Excel: {e}")
        return False

def verify_excel_integrity():
    """Verifica la integridad del archivo Excel."""
    try:
        exe_dir = get_executable_directory()
        excel_file = os.path.join(exe_dir, "unidades_registro.xlsx")
        
        if not os.path.exists(excel_file):
            print("❌ Archivo Excel no existe")
            return False
        
        # Intentar leer el archivo
        df = pd.read_excel(excel_file)
        print(f"✅ Archivo Excel leído correctamente")
        print(f"📊 Columnas: {list(df.columns)}")
        print(f"📈 Filas: {len(df)}")
        
        return True
        
    except Exception as e:
        print(f"❌ Error verificando integridad del archivo Excel: {e}")
        return False

def main():
    """Función principal para verificar y corregir la persistencia del Excel."""
    print("=" * 60)
    print("🔍 VERIFICADOR DE PERSISTENCIA DE EXCEL")
    print("=" * 60)
    
    # Verificar visibilidad del archivo
    print("\n1️⃣ Verificando visibilidad del archivo Excel...")
    if ensure_excel_visibility():
        print("✅ Archivo Excel es visible y accesible")
    else:
        print("❌ Problema con la visibilidad del archivo Excel")
        
        # Intentar copiar desde ubicación temporal
        print("\n2️⃣ Intentando copiar desde ubicación temporal...")
        if copy_excel_from_temp():
            print("✅ Archivo copiado exitosamente")
        else:
            print("❌ No se pudo copiar el archivo")
            
            # Crear archivo nuevo
            print("\n3️⃣ Creando archivo Excel nuevo...")
            exe_dir = get_executable_directory()
            if create_excel_file(exe_dir):
                print("✅ Archivo Excel creado exitosamente")
            else:
                print("❌ No se pudo crear el archivo Excel")
    
    # Verificar integridad
    print("\n4️⃣ Verificando integridad del archivo...")
    if verify_excel_integrity():
        print("✅ Archivo Excel es válido")
    else:
        print("❌ Problema con la integridad del archivo Excel")
    
    print("\n" + "=" * 60)
    print("🎯 VERIFICACIÓN COMPLETADA")
    print("=" * 60)

if __name__ == "__main__":
    main() 