#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script para restaurar el historial completo desde archivos JSON al archivo Excel.
Este script lee todos los archivos JSON en la carpeta data_json y los convierte
al formato Excel manteniendo el historial completo.
"""

import os
import json
import pandas as pd
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
import glob

def restaurar_historial_excel():
    """Restaura el historial completo desde archivos JSON al Excel."""
    
    EXCEL_FILE = "unidades_registro.xlsx"
    JSON_DIR = "data_json"
    
    print("🔄 Iniciando restauración del historial...")
    
    # Verificar que existe la carpeta de archivos JSON
    if not os.path.exists(JSON_DIR):
        print(f"❌ No se encontró la carpeta {JSON_DIR}")
        return False
    
    # Buscar todos los archivos JSON
    json_files = glob.glob(os.path.join(JSON_DIR, "registro_*.json"))
    if not json_files:
        print(f"❌ No se encontraron archivos JSON en {JSON_DIR}")
        return False
    
    print(f"📁 Encontrados {len(json_files)} archivos JSON")
    
    # Ordenar archivos por timestamp (nombre del archivo)
    json_files.sort()
    
    # Lista para almacenar todos los registros
    all_records = []
    
    # Procesar cada archivo JSON
    for json_file in json_files:
        try:
            with open(json_file, 'r', encoding='utf-8') as f:
                data = json.load(f)
            
            timestamp_str = data.get('timestamp', '')
            turno = data.get('turno', '')
            units = data.get('units', [])
            
            # Convertir timestamp a fecha y hora
            try:
                timestamp = datetime.fromisoformat(timestamp_str)
                fecha = timestamp.strftime('%d/%m/%Y')
                hora = timestamp.strftime('%H:%M:%S')
            except:
                # Si no se puede parsear, usar la fecha del archivo
                file_time = os.path.getmtime(json_file)
                timestamp = datetime.fromtimestamp(file_time)
                fecha = timestamp.strftime('%d/%m/%Y')
                hora = timestamp.strftime('%H:%M:%S')
            
            # Procesar cada unidad
            for unit in units:
                alias = unit.get('alias', '')
                km = unit.get('km', '')
                ap = unit.get('ap', '')
                po = unit.get('po', '')
                turnos_seleccionados = unit.get('turnos_seleccionados', [])
                
                # Crear observaciones
                obs_turno = ', '.join(turnos_seleccionados) if turnos_seleccionados else ''
                
                # Calcular observaciones detalladas
                obs_km = calcular_observacion_km(km)
                obs_ap = calcular_observacion_ap(ap)
                obs_po = calcular_observacion_po(po)
                observaciones = f"KM: {obs_km}, AP: {obs_ap}, PO: {obs_po}"
                
                # Agregar registro
                record = [
                    alias,           # 🚓 UNIDADES
                    km,              # 🚔 KM
                    ap,              # 📌 A.P
                    po,              # 📒 P.O
                    turno,           # 🕵️‍♂️ TURNO
                    obs_turno,       # 🌙 OBS-TURNO
                    hora,            # 🕰️ HORA
                    fecha,           # 📅 FECHA
                    observaciones    # ✍ OBSERVACIONES ⛑
                ]
                
                all_records.append(record)
                
        except Exception as e:
            print(f"⚠️ Error procesando {json_file}: {e}")
            continue
    
    if not all_records:
        print("❌ No se pudieron procesar registros válidos")
        return False
    
    print(f"📊 Procesados {len(all_records)} registros totales")
    
    # Crear DataFrame
    columns = [
        '🚓 UNIDADES', '🚔 KM', '📌 A.P', '📒 P.O', 
        '🕵️‍♂️ TURNO', '🌙 OBS-TURNO', '🕰️ HORA', '📅 FECHA', '✍ OBSERVACIONES ⛑'
    ]
    
    df = pd.DataFrame(all_records, columns=columns)
    
    # Ordenar por fecha y hora
    df['FechaHora'] = pd.to_datetime(df['📅 FECHA'] + ' ' + df['🕰️ HORA'], 
                                    format='%d/%m/%Y %H:%M:%S', 
                                    dayfirst=True)
    df = df.sort_values('FechaHora')
    df = df.drop('FechaHora', axis=1)
    
    # Guardar en Excel con formato
    try:
        with pd.ExcelWriter(EXCEL_FILE, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name='Registro de Unidades')
            
            # Obtener el workbook y worksheet
            workbook = writer.book
            worksheet = writer.sheets['Registro de Unidades']
            
            # Aplicar formato a encabezados
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            
            for col in range(1, len(columns) + 1):
                cell = worksheet.cell(row=1, column=col)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center")
        
        print(f"✅ Historial restaurado exitosamente en {EXCEL_FILE}")
        print(f"📈 Total de registros: {len(df)}")
        return True
        
    except Exception as e:
        print(f"❌ Error guardando Excel: {e}")
        return False

def calcular_observacion_km(km):
    """Calcula observación para KM."""
    if not km or not str(km).strip():
        return "⚠️ 90 KM requerido"
    
    try:
        km_int = int(km)
        if km_int < 0:
            return "❌ KM no puede ser negativo"
        
        objetivo_km = 90
        if km_int >= objetivo_km:
            if km_int > 100:
                excedente = km_int - 100
                return f"✅ COMPLETO (+{excedente} KM)"
            return "✅ COMPLETO"
        else:
            faltante = objetivo_km - km_int
            return f"❌ FALTA {faltante} KM"
            
    except ValueError:
        return "❌ KM debe ser numérico"

def calcular_observacion_ap(ap):
    """Calcula observación para AP."""
    if not ap or not str(ap).strip():
        return "⚠️ 230 AP requerido"
    
    try:
        ap_int = int(ap)
        if ap_int < 0:
            return "❌ AP no puede ser negativo"
        
        objetivo_min = 230
        if ap_int >= objetivo_min:
            return "✅ COMPLETO"
        else:
            faltante_min = objetivo_min - ap_int
            faltante_tacticos = (faltante_min + 29) // 30  # Redondear hacia arriba
            return f"❌ FALTA {faltante_min} MIN ({faltante_tacticos}T)"
            
    except ValueError:
        return "❌ AP debe ser numérico"

def calcular_observacion_po(po):
    """Calcula observación para PO."""
    if not po or not str(po).strip():
        return "⚠️ PO requerido"
    
    try:
        po_int = int(po)
        if po_int < 0:
            return "❌ PO no puede ser negativo"
        
        minimo_requerido = 3
        maximo_permitido = 10
        
        if po_int > maximo_permitido:
            return f"❌ PO máximo {maximo_permitido}"
        elif po_int >= minimo_requerido:
            return f"✅ COMPLETO ({po_int} P.O.)"
        else:
            faltante = minimo_requerido - po_int
            return f"❌ FALTAN {faltante} P.O."
            
    except ValueError:
        return "❌ PO debe ser numérico"

def crear_respaldo_actual():
    """Crea un respaldo del archivo Excel actual antes de restaurar."""
    EXCEL_FILE = "unidades_registro.xlsx"
    BACKUP_FILE = f"unidades_registro_backup_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    
    if os.path.exists(EXCEL_FILE):
        try:
            import shutil
            shutil.copy2(EXCEL_FILE, BACKUP_FILE)
            print(f"💾 Respaldo creado: {BACKUP_FILE}")
            return True
        except Exception as e:
            print(f"⚠️ No se pudo crear respaldo: {e}")
            return False
    return True

if __name__ == "__main__":
    print("=" * 60)
    print("🔄 RESTAURADOR DE HISTORIAL COMPLETO")
    print("=" * 60)
    
    # Crear respaldo del archivo actual
    crear_respaldo_actual()
    
    # Restaurar historial
    success = restaurar_historial_excel()
    
    if success:
        print("\n🎉 ¡Restauración completada exitosamente!")
        print("📋 El archivo unidades_registro.xlsx ahora contiene todo el historial.")
    else:
        print("\n❌ La restauración falló. Revisa los errores anteriores.")
    
    print("=" * 60) 